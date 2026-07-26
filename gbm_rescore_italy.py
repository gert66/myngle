"""
gbm_rescore_italy.py — Merge the new Italy GBM score into a GCS run folder
============================================================================
One-off deploy script for the 2026-07-26 Italy scoring redesign (see
Countries/Italy/GBM_MODEL_ITALIE_2026-07-26.md in the Nextcloud workspace for
the full writeup). Reuses rescore_from_gcs.py's GCS I/O plumbing (download
current/, write a new run/ folder, upload) but replaces the scoring logic:
instead of re-running score_company() over scoring_inputs, it merges in a
pre-computed GBM score (company_id -> {score_1_10, tier, gbm_proba}) read
from a local Excel export.

The existing commercial_fit_score_app / commercial_tier_app fields (what the
Company Hub frontend actually reads) are renamed to
commercial_fit_score_app_legacy / commercial_tier_app_legacy the first time
this runs (never overwritten on a repeat run, so the legacy snapshot always
reflects the original v1 model, not an intermediate GBM run). The GBM values
then become the new commercial_fit_score_app / commercial_tier_app.

Like rescore_from_gcs.py, this NEVER touches current/ directly — it writes to
runs/<run_folder>/, and promoting to current/ is a separate, explicit call to
rescore_from_gcs.promote_run_to_current().
"""

from __future__ import annotations

import argparse
import json
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

import rescore_from_gcs as rfg

MODEL_VERSION = "gbm_v1_2026-07-26"

TIER_THRESHOLDS = [
    (9.0, "\U0001F947 Hot"),
    (7.0, "\U0001F948 Warm"),
    (4.0, "\U0001F949 Cool"),
    (0.0, "❄️ Pass"),
]


def tier_for_score(score: float) -> str:
    for threshold, label in TIER_THRESHOLDS:
        if score >= threshold:
            return label
    return TIER_THRESHOLDS[-1][1]


def load_gbm_scores(xlsx_path: str) -> dict:
    """company_id -> {score_1_10, tier, gbm_proba} from the scored Excel export."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        cid = r[idx["company_id"]]
        score = r[idx["score_1_10"]]
        if cid is None or score is None:
            continue
        out[cid] = {
            "score_1_10": round(float(score), 1),
            "tier": tier_for_score(float(score)),
            "gbm_proba": r[idx["gbm_proba"]],
        }
    return out


def merge_scores_into_detail(detail: dict, gbm: dict, now_iso: str) -> dict:
    new_detail = dict(detail)
    # Preserve the ORIGINAL v1 score/tier under _legacy, once — never overwrite
    # an already-set legacy value on a repeat run.
    if "commercial_fit_score_app_legacy" not in new_detail:
        new_detail["commercial_fit_score_app_legacy"] = detail.get("commercial_fit_score_app")
        new_detail["commercial_tier_app_legacy"] = detail.get("commercial_tier_app")

    new_detail["commercial_fit_score_app"] = gbm["score_1_10"]
    new_detail["commercial_tier_app"] = gbm["tier"]
    new_detail["commercial_fit_score_model_version"] = MODEL_VERSION
    new_detail["gbm_rescore_audit"] = {
        "model_version": MODEL_VERSION,
        "rescored_at": now_iso,
        "gbm_proba": gbm["gbm_proba"],
        "previous_score_app": detail.get("commercial_fit_score_app"),
        "previous_tier_app": detail.get("commercial_tier_app"),
    }
    return new_detail


def merge_scores_into_list_item(item: dict, gbm: dict, now_iso: str) -> dict:
    new_item = dict(item)
    if "commercial_fit_score_app_legacy" not in new_item:
        new_item["commercial_fit_score_app_legacy"] = item.get("commercial_fit_score_app")
        new_item["commercial_tier_app_legacy"] = item.get("commercial_tier_app")
    new_item["commercial_fit_score_app"] = gbm["score_1_10"]
    new_item["commercial_tier_app"] = gbm["tier"]
    new_item["commercial_fit_score_model_version"] = MODEL_VERSION
    return new_item


def build_and_upload(
    bucket: str,
    country_folder: str,
    gbm_xlsx: str,
    run_folder: str,
    work_dir: str,
    upload: bool = True,
) -> dict:
    now = datetime.now(timezone.utc)
    now_iso = now.strftime("%Y-%m-%dT%H:%M:%SZ")

    gbm_scores = load_gbm_scores(gbm_xlsx)
    print(f"Geladen: {len(gbm_scores)} GBM-scores uit {gbm_xlsx}")

    staging = Path(work_dir)
    staging.mkdir(parents=True, exist_ok=True)
    current = rfg.download_current_run(bucket, country_folder, staging / "current")
    print(f"Gedownload: {len(current['list_items'])} list items, "
          f"{len(current['detail_files'])} detail-bucketbestanden")

    tier_before = {}
    tier_after = {}
    matched = 0
    unmatched_ids = []

    new_detail_files = {}
    for filename, bucket_dict in current["detail_files"].items():
        new_bucket = {}
        for cid, detail in bucket_dict.items():
            tier_before[detail.get("commercial_tier_app")] = tier_before.get(detail.get("commercial_tier_app"), 0) + 1
            gbm = gbm_scores.get(cid)
            if gbm is None:
                new_bucket[cid] = detail
                unmatched_ids.append(cid)
                tier_after[detail.get("commercial_tier_app")] = tier_after.get(detail.get("commercial_tier_app"), 0) + 1
                continue
            matched += 1
            new_detail = merge_scores_into_detail(detail, gbm, now_iso)
            new_bucket[cid] = new_detail
            tier_after[new_detail["commercial_tier_app"]] = tier_after.get(new_detail["commercial_tier_app"], 0) + 1
        new_detail_files[filename] = new_bucket

    new_list_items = []
    for item in current["list_items"]:
        gbm = gbm_scores.get(item.get("company_id"))
        if gbm is None:
            new_list_items.append(item)
            continue
        new_list_items.append(merge_scores_into_list_item(item, gbm, now_iso))

    manifest = {
        "schema_version": 1,
        "model_version": MODEL_VERSION,
        "generated_at": now_iso,
        "country_folder": country_folder,
        "run_folder": run_folder,
        "source_current_manifest": current.get("manifest"),
        "companies_total": sum(len(b) for b in current["detail_files"].values()),
        "companies_matched_to_gbm_score": matched,
        "companies_unmatched": len(unmatched_ids),
        "unmatched_company_ids_sample": unmatched_ids[:20],
        "tier_distribution_before": tier_before,
        "tier_distribution_after": tier_after,
        "promoted_to_current": False,
    }

    out_dir = staging / "out"
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / rfg.LIST_FILENAME).write_text(
        json.dumps(new_list_items, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    for filename, bucket_dict in new_detail_files.items():
        (out_dir / filename).write_text(
            json.dumps(bucket_dict, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (out_dir / rfg.CURRENT_MANIFEST_FILENAME).write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

    upload_results = []
    if upload:
        upload_results = rfg.upload_rescored_run(out_dir, bucket, country_folder, run_folder)
    manifest["upload_results"] = upload_results
    manifest["local_output_dir"] = str(out_dir)
    return manifest


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--bucket", default=rfg.DEFAULT_GCS_BUCKET)
    ap.add_argument("--country-folder", default="italy")
    ap.add_argument("--gbm-xlsx", required=True)
    ap.add_argument("--run-folder", default=None)
    ap.add_argument("--work-dir", default=None)
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    run_folder = args.run_folder or f"{datetime.now(timezone.utc).strftime('%Y-%m-%d')}_gbm_v1"
    work_dir = args.work_dir or f"./_gbm_rescore_staging_{run_folder}"

    manifest = build_and_upload(
        args.bucket, args.country_folder, args.gbm_xlsx, run_folder, work_dir,
        upload=not args.dry_run,
    )
    print(json.dumps({k: v for k, v in manifest.items() if k not in ("upload_results",)},
                      indent=2, ensure_ascii=False, default=str))
    n_failed = sum(1 for r in manifest.get("upload_results", []) if not r.get("success"))
    if n_failed:
        print(f"\n{n_failed} upload(s) FAILED")
    else:
        print(f"\nGeupload naar gs://{args.bucket}/{args.country_folder}/runs/{run_folder}/"
              if not args.dry_run else "\nDRY RUN - niets geupload")


if __name__ == "__main__":
    main()
